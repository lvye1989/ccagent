# -*- coding: utf-8 -*-
"""《Agent 发展趋势》演示文稿 —— 主题与绘图工具库。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree

# ---------------- 配色 ----------------
BG      = RGBColor.from_string('0A1020')
BG_DARK = RGBColor.from_string('070C18')
CARD    = RGBColor.from_string('131E33')
CARD_2  = RGBColor.from_string('1B2A44')
ACCENT  = RGBColor.from_string('22D3EE')
ACCENT2 = RGBColor.from_string('8B5CF6')
ACCENT3 = RGBColor.from_string('F59E0B')
ACCENT4 = RGBColor.from_string('34D399')
PINK    = RGBColor.from_string('F472B6')
TEXT    = RGBColor.from_string('EAF0F8')
MUTED   = RGBColor.from_string('93A4BC')
DIM     = RGBColor.from_string('6C7F99')
LINE    = RGBColor.from_string('27374F')
GRID    = RGBColor.from_string('22354E')

FONT = '微软雅黑'
SW, SH = 13.333, 7.5
TOTAL_PAGES = 10


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    return prs


def blank(prs):
    return prs.slide_layouts[6]


# ---------------- 文本 ----------------
def _apply(run, size=18, bold=False, color=TEXT, italic=False, font=FONT, spacing=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = font
    rPr = run._r.get_or_add_rPr()
    if spacing is not None:
        rPr.set('spc', str(int(spacing * 100)))
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = etree.SubElement(rPr, qn('a:ea'))
    ea.set('typeface', font)


def tbox(slide, x, y, w, h, items, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = it.get('align', align)
        if it.get('line_spacing'):
            p.line_spacing = it['line_spacing']
        p.space_before = Pt(it.get('space_before', 0))
        p.space_after = Pt(it.get('space_after', 0))
        runs = it.get('runs')
        if runs is None:
            runs = [{'text': it.get('text', '')}]
        for rr in runs:
            r = p.add_run()
            r.text = rr.get('text', '')
            _apply(r,
                   size=rr.get('size', it.get('size', 18)),
                   bold=rr.get('bold', it.get('bold', False)),
                   color=rr.get('color', it.get('color', TEXT)),
                   italic=rr.get('italic', it.get('italic', False)),
                   font=rr.get('font', it.get('font', FONT)),
                   spacing=rr.get('spacing', it.get('spacing')))
    return sh


# ---------------- 图形 ----------------
def shape(slide, kind, x, y, w, h, fill=CARD, line=None, radius=0.08, lw=0.75):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    try:
        sh.shadow.inherit = False
    except Exception:
        pass
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.10)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    return sh


def card(slide, x, y, w, h, fill=CARD, line=LINE, radius=0.07):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, radius=radius)


def rect(slide, x, y, w, h, fill=ACCENT):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=fill, line=None)


def ellipse(slide, x, y, d, fill=ACCENT, line=None):
    return shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=fill, line=line)


def hline(slide, x, y, w, color=LINE, width=0.75):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                    Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    return ln


def vline(slide, x, y, h, color=LINE, width=0.75):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                    Inches(x), Inches(y + h))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    return ln


def arrow(slide, x1, y1, x2, y2, color=ACCENT, width=1.5, head=False):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    if head:
        ln.line._get_or_add_ln().append(etree.fromstring(
            '<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
            ' type="triangle" w="med" len="med"/>'))
    return ln


def center_of(r):
    return (r[0] + r[2] / 2.0, r[1] + r[3] / 2.0)


def edge_point(r, toward):
    cx, cy = center_of(r)
    dx, dy = toward[0] - cx, toward[1] - cy
    if dx == 0 and dy == 0:
        return cx, cy
    hw, hh = r[2] / 2.0, r[3] / 2.0
    sx = hw / abs(dx) if dx else float('inf')
    sy = hh / abs(dy) if dy else float('inf')
    s = min(sx, sy)
    return cx + dx * s, cy + dy * s


def connect(slide, from_rect, to_rect, color=GRID, width=1.5, head=False, gap=0.06):
    a = edge_point(from_rect, center_of(to_rect))
    b = edge_point(to_rect, center_of(from_rect))
    dx, dy = b[0] - a[0], b[1] - a[1]
    d = (dx * dx + dy * dy) ** 0.5 or 1.0
    ux, uy = dx / d, dy / d
    return arrow(slide, a[0] + ux * gap, a[1] + uy * gap,
                 b[0] - ux * gap, b[1] - uy * gap, color=color, width=width, head=head)


# ---------------- 版面 ----------------
def header(slide, kicker, title, idx):
    rect(slide, 0.75, 0.50, 0.07, 0.32, fill=ACCENT)
    tbox(slide, 0.98, 0.48, 11.0, 0.32,
         [{'text': kicker, 'size': 11.5, 'bold': True, 'color': ACCENT, 'spacing': 1.2}])
    tbox(slide, 0.75, 0.86, 11.9, 0.66,
         [{'text': title, 'size': 28, 'bold': True, 'color': TEXT, 'line_spacing': 1.0}])
    hline(slide, 0.75, 1.60, 11.83, LINE, 1.0)
    tbox(slide, 0.75, 6.98, 8.0, 0.3,
         [{'text': 'Agent 发展趋势  ·  技术 / 生态 / 落地观察', 'size': 9, 'color': DIM}])
    tbox(slide, 11.48, 6.98, 1.10, 0.3,
         [{'text': '%02d / %d' % (idx, TOTAL_PAGES), 'size': 9, 'bold': True,
           'color': DIM, 'align': PP_ALIGN.RIGHT}])


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def page(prs, dark=False):
    s = prs.slides.add_slide(blank(prs))
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG_DARK if dark else BG
    return s
